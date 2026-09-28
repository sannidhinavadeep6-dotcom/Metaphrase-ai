import io
import os
import sys
from pathlib import Path
import asyncio
from typing import Optional, List

# Ensure parent directory (for database) and backend directory (for utils) are in sys.path
BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
for p in [str(ROOT_DIR), str(BASE_DIR)]:
    if p not in sys.path:
        sys.path.insert(0, p)

from fastapi import FastAPI, HTTPException, Depends, Header, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import docx

import database
from utils.ai_generator import (
    generate_paraphrase, 
    TONE_PROFILES, 
    SUPPORTED_LANGUAGES
)
from utils.text_metrics import get_detailed_metrics
from utils.sentence_editor import (
    generate_sentence_alternatives,
    get_contextual_synonyms
)
from utils.ai_detector import (
    analyze_ai_probability,
    humanize_text
)
from utils.ocr_parser import extract_text_from_image
from utils.originality import (
    calculate_originality_score,
    generate_citations
)
from utils.doc_parser import parse_document_file
from utils.grammar_checker import analyze_grammar_and_spelling
from utils.plagiarism_checker import analyze_plagiarism
from utils.ai_chat import chat_with_co_pilot

# Initialize Database with WAL & Salted Bcrypt
database.init_db()

app = FastAPI(
    title="Metaphrase AI 3.0 Enterprise Engine",
    description="High-Throughput Neural Text Transformation, Multimodal OCR, Humanizer & Academic Verification API",
    version="3.0.0"
)

# Enable CORS for React SPA, extensions, and local dev
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ----------------- Models -----------------
class LoginRequest(BaseModel):
    email: str
    password: str

class RegisterRequest(BaseModel):
    name: str
    email: str
    password: str

class StatusUpdateRequest(BaseModel):
    status: str  # 'accepted', 'rejected', 'pending'

class ActivityLogRequest(BaseModel):
    email: str
    feature_name: str
    action: str
    details: Optional[str] = ""
    word_count: Optional[int] = 0

class ParaphraseRequest(BaseModel):
    text: str
    tone: Optional[str] = "Simple"
    custom_instruction: Optional[str] = None
    target_language: Optional[str] = "English"
    email: Optional[str] = None

class CreatePersonaRequest(BaseModel):
    email: str
    title: str
    instruction: str
    icon: Optional[str] = "sparkles"

class MetricsRequest(BaseModel):
    original_text: str
    paraphrased_text: str
    email: Optional[str] = None

class SentenceRewriteRequest(BaseModel):
    sentence: str
    full_context: Optional[str] = ""
    tone: Optional[str] = "Fluent"
    email: Optional[str] = None

class SynonymsRequest(BaseModel):
    word: str
    sentence_context: Optional[str] = ""
    email: Optional[str] = None

class AiDetectRequest(BaseModel):
    text: str
    email: Optional[str] = None

class HumanizeRequest(BaseModel):
    text: str
    target_language: Optional[str] = "English"
    email: Optional[str] = None

class OriginalityRequest(BaseModel):
    original_text: str
    transformed_text: str
    email: Optional[str] = None

class CitationRequest(BaseModel):
    title: str
    author: Optional[str] = ""
    year: Optional[str] = ""
    source_url: Optional[str] = ""
    email: Optional[str] = None

class ExportDocxRequest(BaseModel):
    original_text: str
    paraphrased_text: str
    tone: Optional[str] = "Simple"
    target_language: Optional[str] = "English"
    email: Optional[str] = None

class GrammarCheckRequest(BaseModel):
    text: str
    email: Optional[str] = None

class PlagiarismCheckRequest(BaseModel):
    text: str
    email: Optional[str] = None

class ChatMessage(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    messages: List[ChatMessage]
    workspace_text: Optional[str] = ""
    email: Optional[str] = None

class TranslateRequest(BaseModel):
    text: str
    target_language: str
    source_language: Optional[str] = "Auto-detect"
    email: Optional[str] = None

# ----------------- Core Health & Info -----------------
@app.get("/api/health")
@app.get("/healthz")
@app.get("/readyz")
def health_check():
    return {
        "status": "healthy",
        "service": "Metaphrase AI 3.0 Enterprise",
        "security": "bcrypt-salted",
        "database": "sqlite-wal",
        "uptime": "24/7"
    }

@app.get("/api/tones")
def get_tones():
    return {"tones": TONE_PROFILES}

@app.get("/api/languages")
def get_languages():
    return {"languages": SUPPORTED_LANGUAGES}

# ----------------- Authentication Routes -----------------
@app.post("/api/auth/login")
def login(req: LoginRequest):
    user_info = database.verify_login(req.email, req.password)
    if not user_info:
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    
    role, status, name = user_info
    if status == 'pending':
        return {"success": False, "status": "pending", "message": "Account is pending administrator approval."}
    elif status == 'rejected':
        return {"success": False, "status": "rejected", "message": "Account access has been rejected."}
    
    return {
        "success": True,
        "status": status,
        "user": {
            "name": name,
            "email": req.email.strip().lower(),
            "role": role
        }
    }

@app.post("/api/auth/register")
def register(req: RegisterRequest):
    if len(req.password) < 6:
        raise HTTPException(status_code=400, detail="Password must be at least 6 characters.")
    
    success = database.register_user(req.name, req.email, req.password)
    if not success:
        raise HTTPException(status_code=409, detail="An account with this email already exists.")
    
    clean_email = req.email.strip().lower()
    clean_name = req.name.strip()
    
    return {
        "success": True, 
        "message": "Account created and auto-approved! Welcome to Metaphrase AI.",
        "user": {
            "name": clean_name,
            "email": clean_email,
            "role": "user",
            "status": "accepted"
        }
    }

# ----------------- Admin User Management & Analytics -----------------
@app.get("/api/admin/dashboard")
def get_admin_dashboard():
    return database.get_admin_dashboard_stats()

@app.get("/api/admin/users")
def get_all_users_admin():
    users = database.get_all_users_with_stats()
    return {"users": users}

@app.get("/api/admin/users/{email}/progress")
def get_user_progress_admin(email: str):
    progress = database.get_user_progress(email)
    if not progress:
        raise HTTPException(status_code=404, detail="User profile not found.")
    return progress

@app.get("/api/admin/users/{email}/activity")
def get_user_activity_admin(email: str):
    logs = database.get_user_activity_logs(email)
    return {"activity_logs": logs}

@app.put("/api/admin/users/{email}/status")
def update_user_status_admin(email: str, req: StatusUpdateRequest):
    if req.status not in ["accepted", "rejected", "pending"]:
        raise HTTPException(status_code=400, detail="Invalid status specified.")
    database.update_status(email, req.status)
    return {"success": True, "email": email, "status": req.status}

@app.post("/api/activity/log")
def log_user_activity(req: ActivityLogRequest):
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name=req.feature_name,
            action=req.action,
            details=req.details or "",
            word_count=req.word_count or 0
        )
    return {"success": True}

# ----------------- Custom Personas Routes -----------------
@app.get("/api/personas")
def get_personas(email: str):
    if not email:
        raise HTTPException(status_code=400, detail="Email is required.")
    personas = database.get_user_personas(email)
    return {"personas": personas}

@app.post("/api/personas")
def create_persona(req: CreatePersonaRequest):
    if not req.title.strip() or not req.instruction.strip():
        raise HTTPException(status_code=400, detail="Title and instructions are required.")
    
    persona_id = database.create_custom_persona(
        email=req.email,
        title=req.title,
        instruction=req.instruction,
        icon=req.icon or "sparkles"
    )
    return {"success": True, "persona_id": persona_id, "message": "Custom persona saved."}

@app.delete("/api/personas/{persona_id}")
def delete_persona(persona_id: int, email: str):
    if not email:
        raise HTTPException(status_code=400, detail="Email is required.")
    success = database.delete_custom_persona(persona_id, email)
    if not success:
        raise HTTPException(status_code=404, detail="Persona not found or unauthorized.")
    return {"success": True, "message": "Persona deleted."}

# ----------------- AI Text Transformation -----------------
@app.post("/api/paraphrase")
async def paraphrase_text(req: ParaphraseRequest):
    if not req.text or not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty.")
    
    result = await asyncio.to_thread(
        generate_paraphrase, 
        req.text, 
        req.tone, 
        req.custom_instruction, 
        req.target_language
    )
    
    if result == "SERVICE_ERROR" or not result:
        raise HTTPException(status_code=503, detail="AI transformation service currently unavailable. Please try again.")
    
    if req.email:
        record_label = req.tone if not req.custom_instruction else "Custom Persona"
        if req.target_language and req.target_language.lower() != "english":
            record_label += f" ({req.target_language})"
        database.add_history(req.email, req.text, result, record_label)
    
    metrics = get_detailed_metrics(req.text, result)
    ai_detect = analyze_ai_probability(result)
    originality = calculate_originality_score(req.text, result)

    return {
        "paraphrased_text": result,
        "tone": req.tone,
        "target_language": req.target_language or "English",
        "metrics": metrics,
        "ai_detection": ai_detect,
        "originality": originality
    }

# ----------------- Interactive Sentence Editing & Synonyms -----------------
@app.post("/api/sentence/rewrite")
async def rewrite_sentence(req: SentenceRewriteRequest):
    if not req.sentence.strip():
        raise HTTPException(status_code=400, detail="Sentence cannot be empty.")
    
    alternatives = await asyncio.to_thread(
        generate_sentence_alternatives,
        req.sentence,
        req.full_context or "",
        req.tone or "Fluent"
    )
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Sentence Rewriter & Synonyms",
            action="Rewrote sentence alternative",
            details=f"Tone: {req.tone or 'Fluent'} | Sentence: {req.sentence[:60]}...",
            word_count=len(req.sentence.split())
        )
        
    return {"sentence": req.sentence, "alternatives": alternatives}

@app.post("/api/synonyms")
async def get_synonyms(req: SynonymsRequest):
    if not req.word.strip():
        raise HTTPException(status_code=400, detail="Word cannot be empty.")
    
    synonyms = await asyncio.to_thread(
        get_contextual_synonyms,
        req.word,
        req.sentence_context or ""
    )
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Sentence Rewriter & Synonyms",
            action=f"Looked up synonyms for '{req.word}'",
            details=f"Context: {req.sentence_context[:50] if req.sentence_context else ''}",
            word_count=1
        )
        
    return {"word": req.word, "synonyms": synonyms}

# ----------------- AI Detection & Humanizer -----------------
@app.post("/api/ai-detect")
def check_ai_detection(req: AiDetectRequest):
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty.")
    
    result = analyze_ai_probability(req.text)
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="AI Content Detection",
            action="Analyzed AI detection probability",
            details=f"AI Risk: {result.get('ai_probability', 0)}% ({result.get('verdict', '')})",
            word_count=len(req.text.split())
        )
        
    return result

@app.post("/api/humanize")
async def humanize_prose(req: HumanizeRequest):
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty.")
    
    humanized = await asyncio.to_thread(
        humanize_text,
        req.text,
        req.target_language or "English"
    )
    
    if req.email:
        database.add_history(req.email, req.text, humanized, "Humanized")
        
    ai_detect = analyze_ai_probability(humanized)
    metrics = get_detailed_metrics(req.text, humanized)
    
    return {
        "humanized_text": humanized,
        "ai_detection": ai_detect,
        "metrics": metrics
    }

# ----------------- Optical Character Recognition (OCR) -----------------
@app.post("/api/ocr/extract")
async def ocr_extract(
    file: UploadFile = File(...),
    email: Optional[str] = Form(None)
):
    filename = file.filename or "image.png"
    content = await file.read()
    
    if len(content) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Image size exceeds 15MB limit.")
        
    mime_type = file.content_type or "image/png"
    
    try:
        extracted_text = await asyncio.to_thread(
            extract_text_from_image,
            content,
            mime_type
        )
        words_count = len(extracted_text.split())
        
        if email:
            database.log_activity(
                email=email,
                feature_name="Optical Character Recognition (OCR)",
                action=f"Extracted text from image ({filename})",
                details=f"Extracted {words_count} words",
                word_count=words_count
            )
            
        return {
            "filename": filename,
            "extracted_text": extracted_text,
            "word_count": words_count
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

# ----------------- Plagiarism & Citations -----------------
@app.post("/api/originality/check")
def check_originality(req: OriginalityRequest):
    result = calculate_originality_score(req.original_text, req.transformed_text)
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Originality & Plagiarism",
            action="Verified originality & lexical shift",
            details=f"Originality Score: {result.get('originality_score', 0)}% ({result.get('uniqueness_rating', '')})",
            word_count=len(req.original_text.split())
        )
        
    return result

@app.post("/api/citation/generate")
def create_citations(req: CitationRequest):
    result = generate_citations(
        title=req.title,
        author=req.author or "",
        year=req.year or "",
        source_url=req.source_url or ""
    )
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Citation Generator",
            action=f"Generated academic citations for '{req.title}'",
            details=f"Author: {req.author or 'Unknown'} | Year: {req.year or 'N/A'}",
            word_count=0
        )
        
    return result

# ----------------- Universal Document Parser & Batch Processing -----------------
@app.post("/api/doc/extract")
async def extract_doc_text(
    file: UploadFile = File(...),
    email: Optional[str] = Form(None)
):
    filename = file.filename or "document.txt"
    content = await file.read()
    
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File size exceeds 25MB limit.")
        
    try:
        full_text, paragraphs = parse_document_file(filename, content)
        words_count = len(full_text.split())
        
        if email:
            database.log_activity(
                email=email,
                feature_name="Document Parser",
                action=f"Extracted text from document ({filename})",
                details=f"Extracted {words_count} words across {len(paragraphs)} paragraphs",
                word_count=words_count
            )
            
        return {
            "filename": filename,
            "text": full_text,
            "paragraphs_count": len(paragraphs),
            "word_count": words_count
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to extract document text: {e}")

@app.post("/api/batch/upload")
async def upload_batch_document(
    file: UploadFile = File(...),
    tone: Optional[str] = Form("Simple"),
    custom_instruction: Optional[str] = Form(None),
    target_language: Optional[str] = Form("English"),
    email: Optional[str] = Form(None)
):
    filename = file.filename or "document.txt"
    content = await file.read()
    
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File size exceeds 25MB limit.")
    
    try:
        full_text, paragraphs = parse_document_file(filename, content)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse document: {e}")

    if not paragraphs:
        raise HTTPException(status_code=400, detail="No readable text content found in uploaded document.")

    chunks = []
    curr_chunk = []
    curr_words = 0
    for p in paragraphs:
        w_count = len(p.split())
        if curr_words + w_count > 250 and curr_chunk:
            chunks.append("\n\n".join(curr_chunk))
            curr_chunk = [p]
            curr_words = w_count
        else:
            curr_chunk.append(p)
            curr_words += w_count
    if curr_chunk:
        chunks.append("\n\n".join(curr_chunk))

    async def process_chunk(chunk_text):
        return await asyncio.to_thread(
            generate_paraphrase, 
            chunk_text, 
            tone, 
            custom_instruction, 
            target_language
        )

    tasks = [process_chunk(c) for c in chunks]
    transformed_chunks = await asyncio.gather(*tasks)

    compiled_original = "\n\n".join(chunks)
    compiled_paraphrased = "\n\n".join(transformed_chunks)

    if email:
        record_label = f"Batch Doc: {tone}"
        if target_language and target_language.lower() != "english":
            record_label += f" ({target_language})"
        database.add_history(email, f"[{filename}] {compiled_original[:200]}...", compiled_paraphrased[:200] + "...", record_label)

    metrics = get_detailed_metrics(compiled_original, compiled_paraphrased)

    return {
        "filename": filename,
        "total_paragraphs": len(paragraphs),
        "total_chunks": len(chunks),
        "original_text": compiled_original,
        "paraphrased_text": compiled_paraphrased,
        "tone": tone,
        "target_language": target_language,
        "metrics": metrics
    }

# ----------------- Linguistic Metrics Endpoint -----------------
@app.post("/api/metrics")
def calculate_metrics(req: MetricsRequest):
    metrics = get_detailed_metrics(req.original_text, req.paraphrased_text)
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Linguistic Metrics",
            action="Calculated linguistic & readability metrics",
            details=f"Readability: {metrics.get('readability_grade', '')}",
            word_count=len(req.original_text.split())
        )
    return metrics

# ----------------- History Routes -----------------
@app.get("/api/history")
def get_history(email: str):
    if not email:
        raise HTTPException(status_code=400, detail="Email query parameter is required.")
    rows = database.get_user_history(email)
    history = [
        {
            "id": r[0],
            "original_text": r[1],
            "paraphrased_text": r[2],
            "difficulty": r[3],
            "timestamp": r[4]
        }
        for r in rows
    ]
    return {"history": history}

@app.delete("/api/history")
def clear_history(email: str):
    if not email:
        raise HTTPException(status_code=400, detail="Email query parameter is required.")
    database.clear_user_history(email)
    return {"success": True, "message": "History cleared."}

@app.delete("/api/history/{item_id}")
def delete_history_item(item_id: int, email: str):
    if not email:
        raise HTTPException(status_code=400, detail="Email query parameter is required.")
    database.delete_history_item(item_id, email)
    return {"success": True, "message": "Item deleted."}

# ----------------- Document Export Endpoint -----------------
@app.post("/api/export/docx")
def export_docx(req: ExportDocxRequest):
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Microsoft Word Export",
            action="Exported DOCX report",
            details=f"Tone: {req.tone or 'Simple'} | Language: {req.target_language or 'English'}",
            word_count=len(req.paraphrased_text.split())
        )

    doc = docx.Document()
    
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Metaphrase AI 3.0 — Transformation Report")
    title_run.font.bold = True
    title_run.font.size = docx.shared.Pt(18)
    
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Tone Profile: {req.tone} | Language: {req.target_language or 'English'}\n").font.italic = True
    
    h1 = doc.add_heading("Transformed Output", level=1)
    h1.runs[0].font.color.rgb = docx.shared.RGBColor(2, 132, 199)
    doc.add_paragraph(req.paraphrased_text)
    
    doc.add_paragraph("―" * 40)
    
    h2 = doc.add_heading("Original Source Text", level=2)
    h2.runs[0].font.color.rgb = docx.shared.RGBColor(100, 116, 139)
    doc.add_paragraph(req.original_text)
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=Metaphrase_Output.docx"}
    )

# ----------------- Dedicated Grammar Check Endpoint -----------------
@app.post("/api/grammar/check")
async def check_grammar_endpoint(req: GrammarCheckRequest):
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty.")
    
    result = await asyncio.to_thread(analyze_grammar_and_spelling, req.text)
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Grammar Checker",
            action="Analyzed grammar & spelling",
            details=f"Score: {result.get('score', 100)}/100 | Issues: {len(result.get('issues', []))}",
            word_count=result.get("word_count", 0)
        )
        
    return result

# ----------------- Dedicated Plagiarism Check Endpoint -----------------
@app.post("/api/plagiarism/check")
async def check_plagiarism_endpoint(req: PlagiarismCheckRequest):
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty.")
        
    result = await asyncio.to_thread(analyze_plagiarism, req.text)
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="Plagiarism Checker",
            action="Audited plagiarism & web similarity",
            details=f"Originality: {result.get('originality_score', 100)}% ({result.get('verdict', '')})",
            word_count=result.get("word_count", 0)
        )
        
    return result

# ----------------- Dedicated AI Chat Endpoint -----------------
@app.post("/api/chat")
async def chat_endpoint(req: ChatRequest):
    if not req.messages:
        raise HTTPException(status_code=400, detail="Messages list cannot be empty.")
        
    response_text = await asyncio.to_thread(
        chat_with_co_pilot,
        [m.dict() for m in req.messages],
        req.workspace_text or ""
    )
    
    if req.email:
        database.log_activity(
            email=req.email,
            feature_name="AI Chat Co-Pilot",
            action="Chat conversation turn",
            details=f"User prompt: {req.messages[-1].content[:60]}...",
            word_count=len(response_text.split())
        )
        
    return {
        "reply": response_text,
        "role": "assistant"
    }

# ----------------- Dedicated Neural Translation Endpoint -----------------
@app.post("/api/translate")
async def translate_endpoint(req: TranslateRequest):
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty.")
        
    translated = await asyncio.to_thread(
        generate_paraphrase,
        req.text,
        "Fluent",
        f"Translate the text faithfully into {req.target_language} with high natural fluency and proper grammatical syntax.",
        req.target_language
    )
    
    if req.email:
        database.add_history(req.email, req.text, translated, f"Translation ({req.target_language})")
        
    metrics = get_detailed_metrics(req.text, translated)
    
    return {
        "original_text": req.text,
        "translated_text": translated,
        "source_language": req.source_language or "Auto-detect",
        "target_language": req.target_language,
        "metrics": metrics
    }
