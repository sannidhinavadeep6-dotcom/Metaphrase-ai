import markdown
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def markdown_to_word(md_file_path, output_file_path):
    # Read the markdown file
    with open(md_file_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Convert markdown to HTML
    html_content = markdown.markdown(md_content, extensions=['tables', 'fenced_code'])
    
    # Create a new Word document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Parse and add content
    lines = md_content.split('\n')
    
    for line in lines:
        line = line.rstrip()
        
        # Handle headings
        if line.startswith('# '):
            heading = line[2:].strip()
            p = doc.add_heading(heading, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('## '):
            heading = line[3:].strip()
            p = doc.add_heading(heading, level=2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('### '):
            heading = line[4:].strip()
            p = doc.add_heading(heading, level=3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('#### '):
            heading = line[5:].strip()
            p = doc.add_heading(heading, level=4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('##### '):
            heading = line[6:].strip()
            p = doc.add_heading(heading, level=5)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('###### '):
            heading = line[7:].strip()
            p = doc.add_heading(heading, level=6)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Handle bold text
        elif line.startswith('**') and line.endswith('**'):
            bold_text = line[2:-2]
            p = doc.add_paragraph()
            run = p.add_run(bold_text)
            run.bold = True
        
        # Handle code blocks
        elif line.startswith('```'):
            continue  # Skip code block markers
        
        # Handle horizontal rules
        elif line.startswith('---'):
            doc.add_paragraph('_' * 50)
        
        # Handle bullet points
        elif line.startswith('- '):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(bullet_text, style='List Bullet')
        
        # Handle numbered lists
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or \
             line.startswith('4. ') or line.startswith('5. ') or line.startswith('6. ') or \
             line.startswith('7. ') or line.startswith('8. ') or line.startswith('9. ') or \
             line.startswith('10. '):
            num_text = line.split('. ', 1)[1].strip()
            p = doc.add_paragraph(num_text, style='List Number')
        
        # Handle empty lines
        elif not line.strip():
            doc.add_paragraph()
        
        # Handle regular paragraphs
        elif line.strip():
            # Check if it's a keyword line
            if line.startswith('**Keywords:'):
                p = doc.add_paragraph()
                run = p.add_run(line)
                run.bold = True
                run.italic = True
            else:
                doc.add_paragraph(line)
    
    # Save the document
    doc.save(output_file_path)
    print(f"Successfully converted {md_file_path} to {output_file_path}")

if __name__ == "__main__":
    md_file = r"C:\Users\NAVADEEP\Downloads\Metaphrase-ai-main\Metaphrase-ai-main\IEEE_MTech_Documentation.md"
    word_file = r"C:\Users\NAVADEEP\Downloads\Metaphrase-ai-main\Metaphrase-ai-main\IEEE_MTech_Documentation.docx"
    markdown_to_word(md_file, word_file)
